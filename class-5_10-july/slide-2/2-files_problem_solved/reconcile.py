import openpyxl
from openpyxl import Workbook


def load_records(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    records = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] is None:
            continue
        txn_id, date, merchant, category, amount, currency, account = row
        records[txn_id] = {
            "date": date,
            "merchant": merchant,
            "category": category,
            "amount": amount,
            "currency": currency,
            "account": account,
        }
    return records


bank = load_records("Bank_Charges_10_Records.xlsx")
recorded = load_records("Recorded_Spending_10_Records.xlsx")

bank_ids = set(bank)
recorded_ids = set(recorded)

common_ids = bank_ids & recorded_ids
bank_only_ids = bank_ids - recorded_ids
recorded_only_ids = recorded_ids - bank_ids

mismatches = []
for txn_id in sorted(common_ids):
    b, r = bank[txn_id], recorded[txn_id]
    diffs = [f for f in ("date", "merchant", "category", "amount", "account") if b[f] != r[f]]
    if diffs:
        mismatches.append((txn_id, diffs, b, r))

# Fallback fuzzy match: same date + merchant + amount but different ID
bank_only = {tid: bank[tid] for tid in bank_only_ids}
recorded_only = {tid: recorded[tid] for tid in recorded_only_ids}

fuzzy_matches = []
for b_id, b in list(bank_only.items()):
    for r_id, r in list(recorded_only.items()):
        if b["date"] == r["date"] and b["merchant"] == r["merchant"] and b["amount"] == r["amount"]:
            fuzzy_matches.append((b_id, r_id))
            bank_only.pop(b_id, None)
            recorded_only.pop(r_id, None)
            break

print("=== Reconciliation Summary ===")
print(f"Bank records: {len(bank)}, Recorded records: {len(recorded)}")
print(f"Matched by ID: {len(common_ids)}")
print(f"Field mismatches on matched IDs: {len(mismatches)}")
for tid, diffs, b, r in mismatches:
    print(f"  {tid}: differs in {diffs} | bank={b} | recorded={r}")

if fuzzy_matches:
    print(f"\nLikely same transaction under different IDs: {len(fuzzy_matches)}")
    for b_id, r_id in fuzzy_matches:
        print(f"  bank {b_id} <-> recorded {r_id}")

print(f"\nIn Bank but NOT in Recorded ({len(bank_only)}):")
for tid, b in bank_only.items():
    print(f"  {tid}: {b['date'].date()} {b['merchant']} {b['amount']} {b['currency']}")

print(f"\nIn Recorded but NOT in Bank ({len(recorded_only)}):")
for tid, r in recorded_only.items():
    print(f"  {tid}: {r['date'].date()} {r['merchant']} {r['amount']} {r['currency']}")

bank_only_total = sum(b["amount"] for b in bank_only.values())
recorded_only_total = sum(r["amount"] for r in recorded_only.values())
print(f"\nTotal only-in-Bank amount: {bank_only_total:.2f}")
print(f"Total only-in-Recorded amount: {recorded_only_total:.2f}")

# Write report workbook
wb_out = Workbook()
ws1 = wb_out.active
ws1.title = "In Bank Not Recorded"
ws1.append(["Transaction ID", "Date", "Merchant", "Category", "Amount", "Currency", "Account"])
for tid, b in bank_only.items():
    ws1.append([tid, b["date"], b["merchant"], b["category"], b["amount"], b["currency"], b["account"]])

ws2 = wb_out.create_sheet("In Recorded Not Bank")
ws2.append(["Transaction ID", "Date", "Merchant", "Category", "Amount", "Currency", "Account"])
for tid, r in recorded_only.items():
    ws2.append([tid, r["date"], r["merchant"], r["category"], r["amount"], r["currency"], r["account"]])

ws3 = wb_out.create_sheet("Mismatches")
ws3.append(["Transaction ID", "Differing Fields", "Bank Data", "Recorded Data"])
for tid, diffs, b, r in mismatches:
    ws3.append([tid, ", ".join(diffs), str(b), str(r)])

wb_out.save("Reconciliation_Report.xlsx")
print("\nSaved report to Reconciliation_Report.xlsx")
